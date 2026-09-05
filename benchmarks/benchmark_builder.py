"""Synthetic high-cardinality cache benchmark; workbook setup is not timed.

Run: PYTHONPATH=src python benchmarks/benchmark_builder.py --rows 2000 --columns 11
Use the same script and arguments against both revisions for comparisons.
"""
import argparse
import json
from io import BytesIO
import platform
from statistics import median
from time import perf_counter

import openpyxl
from openpyxl_pivots import add_pivot_table


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, default=2000)
    parser.add_argument("--columns", type=int, default=11)
    parser.add_argument("--repeats", type=int, default=3)
    args = parser.parse_args()
    if args.rows < 1 or args.columns < 3 or args.repeats < 1:
        parser.error("rows/repeats must be positive; columns must be at least 3")
    build_times, save_times = [], []
    for _ in range(args.repeats):
        wb = openpyxl.Workbook()
        data = wb.active
        data.title = "Data"
        data.append(["row", "col", "value"] + [f"extra{i}" for i in range(args.columns - 3)])
        for i in range(args.rows):
            data.append([f"R{i % 10}", f"C{i % 4}", i % 100] +
                        [f"unique-{j}-{i}" for j in range(args.columns - 3)])
        target = wb.create_sheet("Pivot")
        start = perf_counter()
        add_pivot_table(target, source=f"Data!A1:{openpyxl.utils.get_column_letter(args.columns)}{args.rows + 1}",
                        destination="A1", name="P", row="row", column="col", value="value")
        build_times.append(perf_counter() - start)
        start = perf_counter()
        wb.save(BytesIO())
        save_times.append(perf_counter() - start)
    print(json.dumps(dict(rows=args.rows, columns=args.columns, repeats=args.repeats,
                          python=platform.python_version(), openpyxl=openpyxl.__version__,
                          build_seconds=median(build_times), save_seconds=median(save_times))))


if __name__ == "__main__":
    main()
