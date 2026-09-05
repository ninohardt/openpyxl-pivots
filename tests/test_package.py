"""Verify the package graph after multiple writes and added PivotTables."""
from io import BytesIO
import posixpath
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl_pivots import PivotBuildError, add_pivot_table, install
from test_regressions import build, NS

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


class PackageTests(unittest.TestCase):
    def test_multiple_pivots_survive_repeated_saves_and_extension(self):
        wb, target, _ = build([("A", "X", 10), (None, "Y", 20)])
        add_pivot_table(target, source="Data!A1:C3", destination="F2", name="Second",
                        row="col", value="value", refresh_on_load=False)
        for cycle in range(3):
            with self.subTest(cycle=cycle):
                stream = BytesIO()
                wb.save(stream)
                with ZipFile(stream) as archive:
                    names = set(archive.namelist())
                    # Every internal relationship must resolve to a real ZIP part.
                    for path in names:
                        if not path.endswith(".rels"):
                            continue
                        root = ET.fromstring(archive.read(path))
                        base = path.rsplit("/_rels/", 1)[0] if "/_rels/" in path else ""
                        ids = []
                        for rel in root:
                            ids.append(rel.get("Id"))
                            if rel.get("TargetMode") == "External":
                                continue
                            destination = rel.get("Target")
                            resolved = destination.lstrip("/") if destination.startswith("/") else posixpath.normpath(posixpath.join(base, destination))
                            self.assertIn(resolved, names, (path, destination))
                        self.assertEqual(len(ids), len(set(ids)))
                    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
                    caches = workbook.findall("x:pivotCaches/x:pivotCache", NS)
                    count = 2 if cycle == 0 else 3
                    self.assertEqual(len(caches), count)
                    cache_ids = {c.get("cacheId") for c in caches}
                    self.assertEqual(len(cache_ids), count)
                    self.assertEqual(len({c.get(RID) for c in caches}), count)
                    for idx in range(1, count + 1):
                        pivot = ET.fromstring(archive.read(f"xl/pivotTables/pivotTable{idx}.xml"))
                        self.assertIn(pivot.get("cacheId"), cache_ids)
                        cache = ET.fromstring(archive.read(f"xl/pivotCache/pivotCacheDefinition{idx}.xml"))
                        fields = cache.find("x:cacheFields", NS)
                        records = ET.fromstring(archive.read(f"xl/pivotCache/pivotCacheRecords{idx}.xml"))
                        self.assertEqual(int(cache.get("recordCount")), len(records))
                        self.assertEqual(int(records.get("count")), len(records))
                        self.assertEqual(int(fields.get("count")), len(fields))
                        for record in records:
                            self.assertEqual(len(record), len(fields))
                            for value, field in zip(record, fields):
                                if value.tag == f"{{{NS['x']}}}x":
                                    self.assertLess(int(value.get("v")), len(field.find("x:sharedItems", NS)))
                stream.seek(0)
                wb = load_workbook(stream)
                self.assertEqual(len(wb["Pivot"]._pivots), 2)
                self.assertFalse(wb["Pivot"]._pivots[1].cache.refreshOnLoad)
                if cycle == 0:
                    add_pivot_table(wb.create_sheet("More"), source="Data!A1:C3", destination="B3",
                                    name="Third", row="row", value="value")

    def test_names_are_unique_across_sheets_case_insensitively(self):
        wb, _, _ = build([("A", "X", 10)])
        with self.assertRaisesRegex(PivotBuildError, "already exists"):
            add_pivot_table(wb.create_sheet("More"), source="Data!A1:C2", destination="A1",
                            name="p", row="row", value="value")

    def test_install_is_idempotent_and_preserves_existing_method(self):
        from openpyxl.worksheet.worksheet import Worksheet
        previous = getattr(Worksheet, "add_pivot_table", None)
        try:
            install()
            method = Worksheet.add_pivot_table
            install()
            self.assertIs(Worksheet.add_pivot_table, method)
            sentinel = object()
            Worksheet.add_pivot_table = sentinel
            install()
            self.assertIs(Worksheet.add_pivot_table, sentinel)
        finally:
            if previous is None:
                del Worksheet.add_pivot_table
            else:
                Worksheet.add_pivot_table = previous
