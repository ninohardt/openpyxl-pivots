"""Public API regressions: values, cache semantics, and failure atomicity."""
from datetime import date, datetime, timezone
from io import BytesIO
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl_pivots import PivotBuildError, add_pivot_table
from openpyxl_pivots.builder import _build_field_cache

NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def build(rows, **kwargs):
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["row", "col", "value"])
    for record in rows:
        data.append(record)
    target = wb.create_sheet("Pivot")
    options = dict(source=f"Data!A1:C{len(rows) + 1}", destination="A1",
                   name="P", row="row", column="col", value="value")
    options.update(kwargs)
    pivot = add_pivot_table(target, **options)
    return wb, target, pivot


def archive_xml(wb, path):
    stream = BytesIO()
    wb.save(stream)
    with ZipFile(stream) as archive:
        return ET.fromstring(archive.read(path))


class CacheRegressionTests(unittest.TestCase):
    def test_blank_cache_matrix_on_both_axes(self):
        for member, flag in [(1, "containsNumber"), (1.5, "containsNumber"),
                             ("North", "containsString"), (True, "containsString"),
                             (date(2025, 1, 1), "containsDate")]:
            for axis in ("row", "col"):
                with self.subTest(member=member, axis=axis):
                    rows = [(member, "X", 10), (None, "X", 20)]
                    if axis == "col":
                        rows = [(c, r, v) for r, c, v in rows]
                    wb, _, _ = build(rows)
                    root = archive_xml(wb, "xl/pivotCache/pivotCacheDefinition1.xml")
                    shared = root.find(f"x:cacheFields/x:cacheField[@name='{axis}']/x:sharedItems", NS)
                    self.assertEqual(shared.get("containsBlank"), "1")
                    self.assertEqual(shared.get("containsSemiMixedTypes"), "1")
                    self.assertEqual(shared.get(flag), "1")
                    self.assertEqual(len(shared.findall("x:m", NS)), 1)
                    self.assertEqual(int(shared.get("count")), len(shared))

    def test_equivalent_numeric_and_date_keys_keep_all_values(self):
        for left, right in [(1, 1.0), (date(2025, 1, 1), datetime(2025, 1, 1))]:
            for column in (None, "col"):
                with self.subTest(left=left, column=column):
                    _, target, pivot = build([(left, left, 10), (right, right, 20)], column=column)
                    self.assertEqual(target.cell(3 if column else 2, 2).value, 30)
                    self.assertEqual(pivot.cache.cacheFields[0].sharedItems.count, 1)
                    if column:
                        self.assertEqual(target["C3"].value, 30)
                        self.assertEqual(target["B4"].value, 30)
                        self.assertEqual(target["C4"].value, 30)

    def test_boolean_numeric_and_text_keys_stay_distinct(self):
        _, target, pivot = build([(True, "X", 10), (1, "X", 20), ("1", "X", 30)], column=None)
        self.assertEqual([target.cell(r, 2).value for r in (2, 3, 4, 5)], [10, 20, 30, 60])
        self.assertEqual(pivot.cache.cacheFields[0].sharedItems.count, 3)

    def test_mixed_date_number_cache_does_not_mix_bounds(self):
        wb, _, _ = build([(date(2025, 1, 1), "X", 10), (2, "X", 20)])
        root = archive_xml(wb, "xl/pivotCache/pivotCacheDefinition1.xml")
        shared = root.find("x:cacheFields/x:cacheField/x:sharedItems", NS)
        self.assertEqual(shared.get("containsMixedTypes"), "1")
        self.assertFalse("minValue" in shared.attrib and "minDate" in shared.attrib)

    def test_long_text_is_declared(self):
        wb, _, _ = build([("a" * 256, "X", 1)])
        root = archive_xml(wb, "xl/pivotCache/pivotCacheDefinition1.xml")
        shared = root.find("x:cacheFields/x:cacheField/x:sharedItems", NS)
        self.assertEqual(shared.get("longText"), "1")

    def test_distinct_cache_values_do_not_require_quadratic_comparisons(self):
        class CountedText(str):
            comparisons = 0
            __hash__ = str.__hash__

            def __eq__(self, other):
                type(self).comparisons += 1
                return str.__eq__(self, other)

        size = 800
        cache = _build_field_cache([CountedText(str(i)) for i in range(size)], categorical=True)
        self.assertEqual(cache.indices, tuple(range(size)))
        self.assertLess(CountedText.comparisons, size * 10)


class AggregationTests(unittest.TestCase):
    def test_all_aggregations_with_blanks_sparse_groups_and_weighted_totals(self):
        rows = [("A", "X", 2), ("A", "X", 4), ("A", "Y", 12),
                ("B", "X", None), ("B", "X", 8)]
        expected = {
            "sum": [[6, 12, 18], [8, 0, 8], [14, 12, 26]],
            "count": [[2, 1, 3], [1, 0, 1], [3, 1, 4]],
            "countNums": [[2, 1, 3], [1, 0, 1], [3, 1, 4]],
            "average": [[3, 12, 6], [8, None, 8], [14 / 3, 12, 6.5]],
            "min": [[2, 12, 2], [8, None, 8], [2, 12, 2]],
            "max": [[4, 12, 12], [8, None, 8], [8, 12, 12]],
        }
        for aggregation, values in expected.items():
            with self.subTest(aggregation=aggregation):
                wb, target, pivot = build(rows, aggregation=aggregation)
                self.assertEqual([[target.cell(r, c).value for c in range(2, 5)] for r in range(3, 6)], values)
                self.assertEqual(pivot.dataFields[0].subtotal, aggregation)
                # Reopening must preserve the materialized values as well as the pivot.
                stream = BytesIO()
                wb.save(stream)
                stream.seek(0)
                loaded = load_workbook(stream)
                self.assertEqual(loaded["Pivot"]["D5"].value, values[-1][-1])

    def test_counts_distinguish_text_boolean_blank_and_numeric_values(self):
        rows = [("A", "X", v) for v in (None, True, "text", 0, 2.5)]
        for agg, total in [("count", 4), ("countNums", 2)]:
            with self.subTest(agg=agg):
                _, target, _ = build(rows, aggregation=agg, column=None)
                self.assertEqual(target["B2"].value, total)


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.data = self.wb.active
        self.data.title = "Data"
        self.data.append(["row", "value"])
        self.data.append(["A", 10])
        self.target = self.wb.create_sheet("Pivot")

    def call(self, **kwargs):
        options = dict(source="Data!A1:B2", destination="A1", name="P", row="row", value="value")
        options.update(kwargs)
        return add_pivot_table(self.target, **options)

    def assert_unchanged_failure(self, **kwargs):
        before = {k: (v.value, v.style_id) for k, v in self.target._cells.items()}
        with self.assertRaises(PivotBuildError):
            self.call(**kwargs)
        after = {k: (v.value, v.style_id) for k, v in self.target._cells.items()}
        self.assertEqual(before, after)
        self.assertEqual(self.target._pivots, [])

    def test_invalid_destinations_fail_before_writes(self):
        for destination in (None, "A", "1", "A0", "A1:B2", "XFE1", "A1048577", "XFD1", "A1048576"):
            with self.subTest(destination=destination):
                self.assert_unchanged_failure(destination=destination)

    def test_invalid_source_ranges(self):
        for source in ("Data!B2:A1", "Data!A0:B2", "Data!A:B", "Data!XFE1:XFF2", (self.data, None)):
            with self.subTest(source=source):
                self.assert_unchanged_failure(source=source)

    def test_source_must_belong_to_same_workbook(self):
        other = Workbook().active
        other.append(["row", "value"])
        other.append(["B", 99])
        self.assert_unchanged_failure(source=(other, "A1:B2"))

    def test_invalid_field_arguments(self):
        for kwargs in ({"row": []}, {"value": []}, {"column": []}, {"column": ""},
                       {"row": "missing"}, {"row": "value"}, {"aggregation": []}):
            with self.subTest(kwargs=kwargs):
                self.assert_unchanged_failure(**kwargs)

    def test_merged_output_rejected_without_partial_write(self):
        self.target.merge_cells("B1:C2")
        self.assert_unchanged_failure()

    def test_occupied_output_rejected_without_partial_write(self):
        self.target["B2"] = "keep me"
        self.assert_unchanged_failure()

    def test_nonfinite_numbers_and_timezone_dates_fail_early(self):
        for value in (float("nan"), float("inf"), float("-inf"), datetime(2025, 1, 1, tzinfo=timezone.utc)):
            with self.subTest(value=value):
                self.data["A2"] = value
                self.assert_unchanged_failure()

    def test_formula_and_error_sources_are_rejected_explicitly(self):
        for value in ("=1+1", "#DIV/0!"):
            with self.subTest(value=value):
                self.data["A2"] = value
                with self.assertRaisesRegex(PivotBuildError, "formulas and Excel errors"):
                    self.call()
                self.assertEqual(self.target._cells, {})

    def test_source_header_validation(self):
        for headers in ((None, "value"), ("row", "row"), ("row", " ")):
            with self.subTest(headers=headers):
                self.data["A1"], self.data["B1"] = headers
                self.assert_unchanged_failure()

    def test_header_only_source(self):
        self.assert_unchanged_failure(source="Data!A1:B1")

    def test_same_sheet_output_cannot_overwrite_blank_source_rows(self):
        self.target = self.data
        # Materialize the blank source row so the snapshot includes it.
        self.data.cell(3, 1)
        self.data.cell(3, 2)
        self.assert_unchanged_failure(source="A1:B3", destination="A3")

    def test_same_sheet_output_outside_source_is_supported(self):
        self.target = self.data
        self.call(source="A1:B2", destination="D1")
        self.assertEqual(self.data["E2"].value, 10)
        self.assertEqual(self.data["B2"].value, 10)

    def test_quoted_sheet_name_ending_in_apostrophe(self):
        self.data.title = "O'Brien'"
        pivot = self.call(source="'O''Brien'''!$A$1:$B$2")
        self.assertEqual(pivot.cache.cacheSource.worksheetSource.sheet, "O'Brien'")

    def test_literal_formula_like_axis_text_stays_text(self):
        self.data["A2"] = "=1+1"
        self.data["A2"].data_type = "s"
        self.call()
        self.assertEqual(self.target["A2"].data_type, "s")
        self.assertEqual(self.target["A2"].value, "=1+1")


if __name__ == "__main__":
    unittest.main()
