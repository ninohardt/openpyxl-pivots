from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from openpyxl_pivots import PivotBuildError, add_pivot_table


def make_workbook():
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["Region", "Year", "Revenue"])
    for row in [
        ("North", 2025, 120),
        ("South", 2025, 80),
        ("North", 2026, 150),
        ("South", 2026, 95),
    ]:
        data.append(row)
    return wb


class BuilderTests(unittest.TestCase):
    def test_creates_complete_package_and_round_trips(self):
        wb = make_workbook()
        target = wb.create_sheet("Pivot")
        pivot = add_pivot_table(
            target,
            source="Data!A1:C5",
            destination="A3",
            name="SalesPivot",
            row="Region",
            column="Year",
            value="Revenue",
        )
        self.assertEqual(pivot.name, "SalesPivot")
        self.assertEqual(target["A3"].value, "Sum of Revenue")
        self.assertEqual(target["B3"].value, "Year")
        self.assertEqual(target["B5"].value, 120)
        self.assertEqual(target["C5"].value, 150)
        self.assertEqual(target["D7"].value, 445)

        with TemporaryDirectory() as directory:
            output = Path(directory) / "pivot.xlsx"
            wb.save(output)
            with ZipFile(output) as archive:
                names = set(archive.namelist())
                self.assertIn("xl/pivotTables/pivotTable1.xml", names)
                self.assertIn("xl/pivotCache/pivotCacheDefinition1.xml", names)
                self.assertIn("xl/pivotCache/pivotCacheRecords1.xml", names)
                self.assertIn(
                    b'refreshOnLoad="1"',
                    archive.read("xl/pivotCache/pivotCacheDefinition1.xml"),
                )
                self.assertIn(
                    b'name="SalesPivot"', archive.read("xl/pivotTables/pivotTable1.xml")
                )
                self.assertIn(
                    b'<location ref="A3:D7" firstHeaderRow="1" firstDataRow="2"',
                    archive.read("xl/pivotTables/pivotTable1.xml"),
                )
                self.assertIn(
                    b'name="SalesPivot" cacheId="1"',
                    archive.read("xl/pivotTables/pivotTable1.xml"),
                )
                pivot_root = ElementTree.fromstring(
                    archive.read("xl/pivotTables/pivotTable1.xml")
                )
                namespace = {
                    "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                }
                first_row_item = pivot_root.find("x:rowItems/x:i", namespace)
                self.assertIsNotNone(first_row_item)
                self.assertEqual(first_row_item.get("t"), "data")
                first_row_index = first_row_item.find("x:x", namespace)
                self.assertIsNotNone(first_row_index)
                self.assertIsNone(first_row_index.get("v"))
                cache_xml = archive.read("xl/pivotCache/pivotCacheDefinition1.xml")
                self.assertNotIn(b'saveData=', cache_xml)
                self.assertIn(b'<cacheField name="Region" numFmtId="0">', cache_xml)

            reloaded = load_workbook(output)
            self.assertEqual(len(reloaded["Pivot"]._pivots), 1)
            self.assertEqual(reloaded["Pivot"]._pivots[0].name, "SalesPivot")
            self.assertEqual(reloaded["Pivot"]._pivots[0].cache.recordCount, 4)

    def test_row_only_average(self):
        wb = make_workbook()
        target = wb.create_sheet("Pivot")
        add_pivot_table(
            target,
            source=(wb["Data"], "A1:C5"),
            destination="B2",
            name="AveragePivot",
            row="Region",
            value="Revenue",
            aggregation="average",
        )
        self.assertEqual(target["C3"].value, 135)
        self.assertEqual(target["C4"].value, 87.5)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "average.xlsx"
            wb.save(output)
            pivot = load_workbook(output)["Pivot"]._pivots[0]
            self.assertEqual(pivot.dataFields[0].subtotal, "average")

    def test_rejects_non_numeric_sum(self):
        wb = make_workbook()
        target = wb.create_sheet("Pivot")
        with self.assertRaisesRegex(PivotBuildError, "requires numeric"):
            add_pivot_table(
                target,
                source="Data!A1:C5",
                destination="A1",
                name="BadPivot",
                row="Year",
                value="Region",
                aggregation="sum",
            )

    def test_text_axis_with_blank_emits_excel_compatible_shared_items(self):
        wb = Workbook()
        data = wb.active
        data.title = "Data"
        data.append(["region", "year", "revenue"])
        data.append(["North", 2025, 120])
        data.append([None, 2025, 80])
        data.append(["North", 2026, 150])

        target = wb.create_sheet("Pivot")
        add_pivot_table(
            target,
            source="Data!A1:C4",
            destination="A3",
            name="BlankAxisPivot",
            row="region",
            column="year",
            value="revenue",
        )

        with TemporaryDirectory() as directory:
            output = Path(directory) / "blank-in-row-field.xlsx"
            wb.save(output)
            with ZipFile(output) as archive:
                root = ElementTree.fromstring(
                    archive.read("xl/pivotCache/pivotCacheDefinition1.xml")
                )
                namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                shared_items = root.find(
                    "x:cacheFields/x:cacheField[@name='region']/x:sharedItems",
                    namespace,
                )
                self.assertIsNotNone(shared_items)
                self.assertEqual(shared_items.get("containsBlank"), "1")
                self.assertEqual(shared_items.get("containsString"), "1")
                self.assertEqual(shared_items.get("containsSemiMixedTypes"), "1")

            pivot = load_workbook(output)["Pivot"]._pivots[0]
            self.assertEqual(pivot.name, "BlankAxisPivot")

    def test_date_axis_with_blank_emits_both_required_flags(self):
        wb = Workbook()
        data = wb.active
        data.title = "Data"
        data.append(["day", "region", "revenue"])
        data.append([date(2025, 1, 1), "North", 120])
        data.append([None, "South", 80])
        data.append([date(2025, 1, 2), "North", 150])

        target = wb.create_sheet("Pivot")
        add_pivot_table(
            target,
            source="Data!A1:C4",
            destination="A3",
            name="BlankDatePivot",
            row="day",
            column="region",
            value="revenue",
        )

        with TemporaryDirectory() as directory:
            output = Path(directory) / "blank-in-date-field.xlsx"
            wb.save(output)
            with ZipFile(output) as archive:
                root = ElementTree.fromstring(
                    archive.read("xl/pivotCache/pivotCacheDefinition1.xml")
                )
                namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                shared_items = root.find(
                    "x:cacheFields/x:cacheField[@name='day']/x:sharedItems",
                    namespace,
                )
                self.assertIsNotNone(shared_items)
                self.assertEqual(shared_items.get("containsBlank"), "1")
                self.assertEqual(shared_items.get("containsDate"), "1")
                self.assertEqual(shared_items.get("containsSemiMixedTypes"), "1")

    def test_boolean_axis_with_blank_marks_items_as_strings(self):
        wb = Workbook()
        data = wb.active
        data.title = "Data"
        data.append(["active", "region", "revenue"])
        data.append([True, "North", 120])
        data.append([None, "South", 80])
        data.append([False, "North", 150])

        target = wb.create_sheet("Pivot")
        add_pivot_table(
            target,
            source="Data!A1:C4",
            destination="A3",
            name="BlankBooleanPivot",
            row="active",
            column="region",
            value="revenue",
        )

        with TemporaryDirectory() as directory:
            output = Path(directory) / "blank-in-boolean-field.xlsx"
            wb.save(output)
            with ZipFile(output) as archive:
                root = ElementTree.fromstring(
                    archive.read("xl/pivotCache/pivotCacheDefinition1.xml")
                )
                namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                shared_items = root.find(
                    "x:cacheFields/x:cacheField[@name='active']/x:sharedItems",
                    namespace,
                )
                self.assertIsNotNone(shared_items)
                self.assertEqual(shared_items.get("containsBlank"), "1")
                self.assertEqual(shared_items.get("containsSemiMixedTypes"), "1")
                self.assertEqual(shared_items.get("containsString"), "1")

    def test_rejects_multiple_value_fields_explicitly(self):
        wb = make_workbook()
        target = wb.create_sheet("Pivot")
        with self.assertRaisesRegex(
            PivotBuildError,
            "value must be a single source field name",
        ):
            add_pivot_table(
                target,
                source="Data!A1:C5",
                destination="A1",
                name="MultipleValues",
                row="Region",
                column="Year",
                value=["Revenue", "Year"],
            )


if __name__ == "__main__":
    unittest.main()
