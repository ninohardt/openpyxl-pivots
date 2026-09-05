"""Build a small, worksheet-backed OOXML PivotTable object graph.

openpyxl already contains serializers and package writers for PivotTables.  It
does not expose a creation API, however.  This module derives the required
cache and table definitions from a source range and attaches them to the
target worksheet using openpyxl's existing writer pipeline.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from math import isfinite
from numbers import Real
from typing import Any, Iterable

from openpyxl.pivot.cache import (
    CacheDefinition,
    CacheField,
    CacheSource,
    SharedItems,
    WorksheetSource,
)
from openpyxl.pivot.fields import (
    Boolean,
    DateTimeField,
    Index,
    Missing,
    Number,
    Text,
)
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import (
    DataField,
    FieldItem,
    Location,
    PivotField,
    PivotTableStyle,
    RowColField,
    RowColItem,
    TableDefinition,
)
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.cell import (
    get_column_letter,
    range_boundaries,
)


class PivotBuildError(ValueError):
    """Raised when a requested PivotTable cannot be represented by the MVP."""


@dataclass(frozen=True)
class _Source:
    worksheet: Any
    ref: str
    min_col: int
    min_row: int
    max_col: int
    max_row: int


@dataclass(frozen=True)
class _FieldCache:
    shared_items: SharedItems
    indices: tuple[int, ...] | None


_AGGREGATION_LABELS = {
    "sum": "Sum",
    "count": "Count",
    "countNums": "Count",
    "average": "Average",
    "min": "Min",
    "max": "Max",
}


def add_pivot_table(
    worksheet,
    *,
    source: str | tuple[Any, str],
    destination: str,
    name: str,
    row: str,
    value: str,
    column: str | None = None,
    aggregation: str = "sum",
    style: str = "PivotStyleMedium9",
    refresh_on_load: bool = True,
):
    """Create and attach a worksheet-backed PivotTable.

    Parameters are deliberately narrow while creation support is experimental.
    ``source`` may be ``"Data!A1:D100"``, ``"A1:D100"`` (same worksheet), or
    ``(source_worksheet, "A1:D100")``.
    """
    if not isinstance(worksheet, Worksheet):
        raise PivotBuildError("target must be a normal, writable worksheet")
    if not isinstance(aggregation, str) or aggregation not in _AGGREGATION_LABELS:
        allowed = ", ".join(_AGGREGATION_LABELS)
        raise PivotBuildError(f"aggregation must be one of: {allowed}")

    wb = worksheet.parent
    field_arguments = {"row": row, "value": value}
    if column is not None:
        field_arguments["column"] = column
    for argument, field in field_arguments.items():
        if not isinstance(field, str) or not field:
            raise PivotBuildError(
                f"{argument} must be a single source field name, got {type(field).__name__}"
            )

    if not isinstance(style, str) or not style:
        raise PivotBuildError("style must be a non-empty string")
    if not isinstance(refresh_on_load, bool):
        raise PivotBuildError("refresh_on_load must be a boolean")
    _validate_name(wb, name)
    _destination_cell(destination)
    source_info = _resolve_source(worksheet, source)
    headers, records = _read_source(source_info)
    field_index = {header: idx for idx, header in enumerate(headers)}
    requested = [row, value] + ([column] if column is not None else [])
    missing = [field for field in requested if field not in field_index]
    if missing:
        raise PivotBuildError(f"unknown source field(s): {', '.join(missing)}")
    if len(set(requested)) != len(requested):
        raise PivotBuildError("row, column, and value fields must be different in this MVP")
    if not records:
        raise PivotBuildError("the source range has headers but no data rows")

    row_idx = field_index[row]
    value_idx = field_index[value]
    col_idx = field_index[column] if column else None
    _validate_values(records, value_idx, aggregation, value)

    field_caches = [
        _build_field_cache([record[idx] for record in records], categorical=(idx != value_idx))
        for idx in range(len(headers))
    ]
    cache_fields = [
        CacheField(
            name=header,
            sharedItems=field_cache.shared_items,
            uniqueList=None,
            numFmtId=0,
            sqlType=None,
            hierarchy=None,
            level=None,
            databaseField=None,
        )
        for header, field_cache in zip(headers, field_caches)
    ]
    cache_records = _build_cache_records(records, field_caches, value_idx)

    cache_id = _next_cache_id(wb)
    cache = CacheDefinition(
        saveData=None,
        # openpyxl deduplicates caches by serialized definition equality, but
        # does not repair the referencing cacheIds. Persist per-pivot builder
        # provenance so independent caches stay distinct after save/load/save.
        refreshedBy=f"openpyxl-pivots ({name})",
        refreshOnLoad=refresh_on_load,
        enableRefresh=None,
        createdVersion=3,
        refreshedVersion=8,
        minRefreshableVersion=3,
        recordCount=len(records),
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(
                ref=source_info.ref,
                sheet=source_info.worksheet.title,
            ),
        ),
        cacheFields=cache_fields,
    )
    cache.records = RecordList(r=cache_records)

    row_keys = _ordered_unique(record[row_idx] for record in records)
    col_keys = _ordered_unique(record[col_idx] for record in records) if column else []
    if worksheet is source_info.worksheet:
        left, top = _destination_cell(destination)
        right = left + (len(col_keys) + 2 if column else 2) - 1
        bottom = top + len(row_keys) + (3 if column else 2) - 1
        if not (right < source_info.min_col or left > source_info.max_col
                or bottom < source_info.min_row or top > source_info.max_row):
            raise PivotBuildError("PivotTable output must not overlap its source range")
    result_ref = _render_result(
        worksheet,
        destination,
        row,
        column,
        value,
        aggregation,
        records,
        row_idx,
        col_idx,
        value_idx,
        row_keys,
        col_keys,
    )

    pivot_fields = []
    for idx, field_cache in enumerate(field_caches):
        axis = None
        data_field = None
        items = ()
        if idx == row_idx:
            axis = "axisRow"
            items = _field_items(field_cache)
        elif idx == col_idx:
            axis = "axisCol"
            items = _field_items(field_cache)
        elif idx == value_idx:
            data_field = True
        pivot_fields.append(
            PivotField(
                axis=axis,
                dataField=data_field,
                items=items,
                showAll=False,
                compact=False,
                outline=False,
                subtotalTop=False,
                defaultSubtotal=True,
                includeNewItemsInFilter=True,
            )
        )

    row_item_indices = _key_indices(field_caches[row_idx], row_keys)
    row_items = _axis_items(row_item_indices)

    if column:
        col_item_indices = _key_indices(field_caches[col_idx], col_keys)
        col_fields = [RowColField(x=col_idx)]
        col_items = _axis_items(col_item_indices)
    else:
        col_fields = []
        col_items = [RowColItem()]

    first_data_row = 2 if column else 1
    pivot = TableDefinition(
        name=name,
        cacheId=cache_id,
        dataOnRows=False,
        dataCaption="Values",
        updatedVersion=8,
        minRefreshableVersion=3,
        createdVersion=3,
        useAutoFormatting=True,
        applyWidthHeightFormats=True,
        compact=False,
        compactData=False,
        outline=False,
        gridDropZones=True,
        multipleFieldFilters=False,
        location=Location(
            ref=result_ref,
            firstHeaderRow=1,
            firstDataRow=first_data_row,
            firstDataCol=1,
        ),
        pivotFields=pivot_fields,
        rowFields=[RowColField(x=row_idx)],
        rowItems=row_items,
        colFields=col_fields,
        colItems=col_items,
        dataFields=[
            DataField(
                name=f"{_AGGREGATION_LABELS[aggregation]} of {value}",
                fld=value_idx,
                subtotal=aggregation,
                baseField=0,
                baseItem=0,
            )
        ],
        pivotTableStyleInfo=PivotTableStyle(
            name=style,
            showRowHeaders=True,
            showColHeaders=True,
            showRowStripes=False,
            showColStripes=False,
            showLastColumn=False,
        ),
    )
    pivot.cache = cache
    worksheet.add_pivot(pivot)
    return pivot


def _resolve_source(target_ws, source) -> _Source:
    if isinstance(source, tuple):
        if len(source) != 2:
            raise PivotBuildError("source tuple must be (worksheet, range)")
        source_ws, ref = source
    elif isinstance(source, str):
        if "!" in source:
            sheet_token, ref = source.rsplit("!", 1)
            if sheet_token.startswith("'") and sheet_token.endswith("'"):
                sheet_name = sheet_token[1:-1].replace("''", "'")
            else:
                sheet_name = sheet_token
            try:
                source_ws = target_ws.parent[sheet_name]
            except KeyError as exc:
                raise PivotBuildError(f"source worksheet does not exist: {sheet_name}") from exc
        else:
            source_ws, ref = target_ws, source
    else:
        raise PivotBuildError("source must be a range string or (worksheet, range)")

    if not isinstance(source_ws, Worksheet) or source_ws.parent is not target_ws.parent:
        raise PivotBuildError("source worksheet must belong to the same workbook as the target")
    if not isinstance(ref, str):
        raise PivotBuildError("source range must be a string")
    ref = ref.replace("$", "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
    except ValueError as exc:
        raise PivotBuildError(f"invalid source range: {ref}") from exc
    if None in (min_col, min_row, max_col, max_row):
        raise PivotBuildError("source must be a bounded rectangular range")
    if not (1 <= min_col <= max_col <= 16384 and 1 <= min_row <= max_row <= 1048576):
        raise PivotBuildError("source range must be ordered and within Excel worksheet limits")
    return _Source(source_ws, ref, min_col, min_row, max_col, max_row)


def _read_source(source: _Source) -> tuple[list[str], list[list[Any]]]:
    rows = list(
        source.worksheet.iter_rows(
            min_row=source.min_row,
            max_row=source.max_row,
            min_col=source.min_col,
            max_col=source.max_col,
            values_only=False,
        )
    )
    if not rows:
        raise PivotBuildError("source range is empty")
    raw_headers = [cell.value for cell in rows[0]]
    if any(value is None or str(value).strip() == "" for value in raw_headers):
        raise PivotBuildError("every source column must have a non-empty header")
    headers = [str(value) for value in raw_headers]
    if len(set(headers)) != len(headers):
        raise PivotBuildError("source headers must be unique")
    records = []
    for cells in rows[1:]:
        record = []
        for cell in cells:
            value = cell.value
            location = f"{source.worksheet.title}!{cell.coordinate}"
            if cell.data_type in {"f", "e"}:
                raise PivotBuildError(f"formulas and Excel errors are unsupported in source cell {location}; use literal or cached values")
            if _is_number(value) and not isfinite(value):
                raise PivotBuildError(f"source cell {location} must contain a finite number")
            if isinstance(value, datetime) and value.tzinfo is not None:
                raise PivotBuildError(f"source cell {location} must contain a timezone-naive datetime")
            if value is not None and not isinstance(value, (str, bool, date)) and not _is_number(value):
                raise PivotBuildError(f"unsupported source value type in {location}: {type(value).__name__}")
            record.append(value)
        records.append(record)
    return headers, records


def _validate_name(workbook, name: str) -> None:
    if not isinstance(name, str) or not name.strip():
        raise PivotBuildError("name must be a non-empty string")
    existing = {
        pivot.name.casefold()
        for ws in workbook.worksheets
        for pivot in ws._pivots
    }
    if name.casefold() in existing:
        raise PivotBuildError(f"PivotTable name already exists: {name}")


def _validate_values(records, value_idx, aggregation, field_name):
    if aggregation in {"count"}:
        return
    nonblank = [row[value_idx] for row in records if row[value_idx] is not None]
    if aggregation == "countNums":
        return
    invalid = [value for value in nonblank if not _is_number(value)]
    if invalid:
        raise PivotBuildError(
            f"{aggregation} requires numeric values in {field_name!r}; "
            f"found {invalid[0]!r}"
        )


def _is_number(value: Any) -> bool:
    return isinstance(value, Real) and not isinstance(value, bool)


def _ordered_unique(values: Iterable[Any]) -> list[Any]:
    result = {}
    for value in values:
        result.setdefault(_cache_key(value), value)
    return list(result.values())


def _cache_key(value: Any) -> tuple:
    """Use the same OOXML identity for cache indices and displayed groups.

    Numbers share one category, as do dates and datetimes. Booleans remain
    distinct from numbers even though Python considers True == 1.
    """
    if value is None:
        return ("m", None)
    if isinstance(value, bool):
        return ("b", value)
    if _is_number(value):
        return ("n", float(value))
    if isinstance(value, (date, datetime)):
        return ("d", _as_datetime(value))
    return ("s", value)


def _build_field_cache(values: list[Any], *, categorical: bool) -> _FieldCache:
    if not categorical and all(value is None or _is_number(value) for value in values):
        numbers = [float(value) for value in values if value is not None]
        return _FieldCache(
            SharedItems(
                _fields=(),
                containsInteger=bool(numbers) and all(number.is_integer() for number in numbers),
                containsNumber=bool(numbers),
                containsSemiMixedTypes=False,
                containsString=False,
                minValue=min(numbers) if numbers else None,
                maxValue=max(numbers) if numbers else None,
            ),
            None,
        )

    unique = _ordered_unique(values)
    fields = [_pivot_value(value) for value in unique]
    lookup = {_cache_key(value): idx for idx, value in enumerate(unique)}
    indices = tuple(lookup[_cache_key(value)] for value in values)
    types = {_cache_key(value)[0] for value in unique if value is not None}
    has_blank = any(value is None for value in unique)
    numeric = [float(value) for value in unique if _is_number(value)]
    dates = [_as_datetime(value) for value in unique if isinstance(value, (date, datetime))]
    shared_kwargs = {"_fields": fields}
    # Excel emits no redundant type flags for a pure string cache. For numeric
    # caches it emits only the numeric characteristics below. Keeping the XML
    # close to Excel's own output avoids a cache-repair pass on open.
    if types == {"n"} and not has_blank:
        shared_kwargs.update(
            containsInteger=bool(numeric) and all(number.is_integer() for number in numeric),
            containsNumber=True,
            containsSemiMixedTypes=False,
            containsString=False,
            minValue=min(numeric) if numeric else None,
            maxValue=max(numeric) if numeric else None,
        )
    elif types == {"d"} and not has_blank:
        shared_kwargs.update(
            containsDate=True,
            containsNonDate=False,
            containsSemiMixedTypes=False,
            containsString=False,
            minDate=min(dates),
            maxDate=max(dates),
        )
    elif len(types) > 1 or has_blank:
        shared_kwargs.update(
            containsBlank=has_blank,
            containsDate=bool(dates),
            containsInteger=bool(numeric) and all(number.is_integer() for number in numeric),
            containsMixedTypes=len(types) > 1,
            containsNonDate=any(
                not isinstance(value, (date, datetime))
                for value in unique
                if value is not None
            ),
            containsNumber=bool(numeric),
            # Office requires this for text, blank, boolean, or error items.
            containsSemiMixedTypes=any(
                value is None or isinstance(value, (str, bool)) for value in unique
            ),
            containsString=any(isinstance(value, (str, bool)) for value in unique),
            # OOXML date bounds must not be mixed with numeric bounds.
            minValue=min(numeric) if numeric and not dates else None,
            maxValue=max(numeric) if numeric and not dates else None,
            minDate=min(dates) if dates and not numeric else None,
            maxDate=max(dates) if dates and not numeric else None,
        )
    if any(isinstance(value, str) and len(value) > 255 for value in unique):
        shared_kwargs["longText"] = True
    return _FieldCache(SharedItems(**shared_kwargs), indices)


def _pivot_value(value: Any):
    if value is None:
        return Missing()
    if isinstance(value, bool):
        return Boolean(v=value)
    if isinstance(value, datetime):
        return DateTimeField(v=value)
    if isinstance(value, date):
        return DateTimeField(v=_as_datetime(value))
    if _is_number(value):
        return Number(v=float(value))
    return Text(v=str(value))


def _as_datetime(value: date | datetime) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime(value.year, value.month, value.day)


def _build_cache_records(records, field_caches, value_idx):
    output = []
    for row_number, record in enumerate(records):
        fields = []
        for idx, value in enumerate(record):
            cache = field_caches[idx]
            if cache.indices is not None:
                fields.append(Index(v=cache.indices[row_number]))
            else:
                fields.append(_pivot_value(value))
        output.append(Record(_fields=fields))
    return output


def _field_items(field_cache: _FieldCache):
    if field_cache.indices is None:
        return [FieldItem(t="default")]
    count = field_cache.shared_items.count
    return [FieldItem(x=idx) for idx in range(count)] + [FieldItem(t="default")]


def _axis_items(indices: list[int]) -> list[RowColItem]:
    """Encode axis item zero using OOXML's omitted-value default.

    Excel writes the first member as ``<i><x/></i>`` and grand totals as
    ``<i t="grand"><x/></i>``.  A literal ``v="0"`` is semantically similar
    but was one of the structures Excel normalized while repairing v0.1.
    """
    items = [
        RowColItem(x=[Index(v=None if idx == 0 else idx)])
        for idx in indices
    ]
    items.append(RowColItem(t="grand", x=[Index(v=None)]))
    return items


def _key_indices(field_cache: _FieldCache, keys: list[Any]) -> list[int]:
    if field_cache.indices is None:
        raise AssertionError("axis field did not receive a shared-item cache")
    values = [_value_from_pivot(item) for item in field_cache.shared_items._fields]
    lookup = {_cache_key(value): idx for idx, value in enumerate(values)}
    return [lookup[_cache_key(key)] for key in keys]


def _value_from_pivot(item):
    if isinstance(item, Missing):
        return None
    return item.v


def _next_cache_id(workbook) -> int:
    ids = [
        pivot.cacheId
        for ws in workbook.worksheets
        for pivot in ws._pivots
        if pivot.cacheId is not None
    ]
    # Zero is schema-valid, but native Excel-created caches conventionally use
    # positive identifiers and Excel reassigned the v0.1 cache during repair.
    return max(ids, default=0) + 1


def _aggregate(values: list[Any], aggregation: str):
    present = [value for value in values if value is not None]
    if aggregation == "count":
        return len(present)
    if aggregation == "countNums":
        return sum(_is_number(value) for value in present)
    numeric = [float(value) for value in present]
    if aggregation == "sum":
        return sum(numeric)
    if not numeric:
        return None
    if aggregation == "average":
        return sum(numeric) / len(numeric)
    if aggregation == "min":
        return min(numeric)
    if aggregation == "max":
        return max(numeric)
    raise AssertionError(f"unhandled aggregation {aggregation}")


def _render_result(
    ws,
    destination,
    row_name,
    column_name,
    value_name,
    aggregation,
    records,
    row_idx,
    col_idx,
    value_idx,
    row_keys,
    col_keys,
):
    min_col, min_row = _destination_cell(destination)

    groups = defaultdict(list)
    column_values = defaultdict(list)
    for record in records:
        key = (record[row_idx], record[col_idx] if col_idx is not None else None)
        groups[_hashable_key(key)].append(record[value_idx])
        if col_idx is not None:
            column_values[_cache_key(record[col_idx])].append(record[value_idx])

    height = len(row_keys) + (3 if column_name else 2)
    width = (len(col_keys) + 2) if column_name else 2
    _ensure_blank(ws, min_row, min_col, height, width)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    bold = Font(bold=True)

    if column_name:
        ws.cell(
            min_row,
            min_col,
            f"{_AGGREGATION_LABELS[aggregation]} of {value_name}",
        )
        _write_label(ws, min_row, min_col + 1, column_name)
        label_row = min_row + 1
        _write_label(ws, label_row, min_col, row_name)
        for offset, key in enumerate(col_keys, 1):
            _write_label(ws, label_row, min_col + offset, _display_key(key))
        ws.cell(label_row, min_col + len(col_keys) + 1, "Grand Total")
    else:
        label_row = min_row
        _write_label(ws, min_row, min_col, row_name)
        ws.cell(
            min_row,
            min_col + 1,
            f"{_AGGREGATION_LABELS[aggregation]} of {value_name}",
        )

    for header_row in range(min_row, label_row + 1):
        for col in range(min_col, min_col + width):
            ws.cell(header_row, col).font = bold
            ws.cell(header_row, col).fill = header_fill

    for row_offset, row_key in enumerate(row_keys, 1):
        out_row = label_row + row_offset
        _write_label(ws, out_row, min_col, _display_key(row_key))
        if column_name:
            all_values = []
            for col_offset, col_key in enumerate(col_keys, 1):
                values = groups.get(_hashable_key((row_key, col_key)), [])
                all_values.extend(values)
                ws.cell(out_row, min_col + col_offset, _aggregate(values, aggregation))
            ws.cell(out_row, min_col + len(col_keys) + 1, _aggregate(all_values, aggregation))
        else:
            values = groups.get(_hashable_key((row_key, None)), [])
            ws.cell(out_row, min_col + 1, _aggregate(values, aggregation))

    total_row = label_row + len(row_keys) + 1
    ws.cell(total_row, min_col, "Grand Total")
    if column_name:
        all_values = []
        for col_offset, col_key in enumerate(col_keys, 1):
            values = column_values[_cache_key(col_key)]
            all_values.extend(values)
            ws.cell(total_row, min_col + col_offset, _aggregate(values, aggregation))
        ws.cell(total_row, min_col + len(col_keys) + 1, _aggregate(all_values, aggregation))
    else:
        ws.cell(total_row, min_col + 1, _aggregate([r[value_idx] for r in records], aggregation))

    for col in range(min_col, min_col + width):
        ws.cell(total_row, col).font = bold
        ws.cell(total_row, col).fill = total_fill

    # Pivot labels such as "Sum of Revenue" and "Grand Total" should be
    # visible before Excel performs its own refresh/autofit pass.
    for col in range(min_col, min_col + width):
        contents = [ws.cell(row, col).value for row in range(min_row, total_row + 1)]
        natural_width = max((len(str(value)) for value in contents if value is not None), default=8) + 2
        ws.column_dimensions[get_column_letter(col)].width = min(max(natural_width, 10), 22)

    return (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(min_col + width - 1)}{total_row}"
    )


def _hashable_key(key):
    return tuple(_cache_key(value) for value in key)


def _write_label(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"


def _destination_cell(destination):
    if not isinstance(destination, str):
        raise PivotBuildError("destination must be a single cell")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(destination.replace("$", ""))
    except ValueError as exc:
        raise PivotBuildError(f"invalid destination: {destination}") from exc
    if (None in (min_col, min_row, max_col, max_row)
            or (min_col, min_row) != (max_col, max_row)
            or not (1 <= min_col <= 16384 and 1 <= min_row <= 1048576)):
        raise PivotBuildError("destination must be a single cell within Excel worksheet limits")
    return min_col, min_row


def _display_key(value):
    return "(blank)" if value is None else value


def _ensure_blank(ws, min_row, min_col, height, width):
    max_row, max_col = min_row + height - 1, min_col + width - 1
    if max_row > 1048576 or max_col > 16384:
        raise PivotBuildError("PivotTable output exceeds Excel worksheet limits")
    for merged in ws.merged_cells.ranges:
        if not (max_col < merged.min_col or min_col > merged.max_col
                or max_row < merged.min_row or min_row > merged.max_row):
            raise PivotBuildError(f"PivotTable output overlaps merged cells: {merged}")
    occupied = []
    for row in range(min_row, min_row + height):
        for col in range(min_col, min_col + width):
            cell = ws._cells.get((row, col))
            if cell is not None and cell.value is not None:
                occupied.append(cell.coordinate)
    if occupied:
        raise PivotBuildError(
            "PivotTable output would overwrite non-empty cells: " + ", ".join(occupied[:5])
        )
